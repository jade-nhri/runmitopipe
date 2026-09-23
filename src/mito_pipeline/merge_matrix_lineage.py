#!/usr/bin/env python3
"""
merge_matrix_lineage.py

Left-joins genome_primer_matrix.tsv (from evaluate_primers.py) with
genome_lineage.tsv (from build_taxonomy_summary.py, run with
--all-genomes) on genome_id, and writes the combined table to an Excel
(.xlsx) file. With both upstream scripts' default settings, the resulting
column layout matches Supplementary Table S1/S2:

    genome_id, Primer1 ... PrimerN, n_primers_hit, has_plus_strand,
    has_minus_strand, detectable, taxid, kingdom, phylum, class, order,
    family, genus, species

  - genome_primer_matrix.tsv is the LEFT/primary table, so every row in
    it is kept (including detectable == 0 genomes, and genomes with no
    taxid match, which simply show blank lineage columns).
  - genome_lineage.tsv columns (taxid + the 7 rank columns) are appended
    on the right.

If --out already exists, the sheet is added to (or replaces a
same-named sheet in) that workbook rather than overwriting the whole
file -- so running this once per dataset with --sheet-name TableS1 and
then --sheet-name TableS2 against the same --out produces one workbook
with both sheets, matching Supplementary Tables.xlsx.

Usage (produces one workbook with both sheets):
    python3 merge_matrix_lineage.py \
        --matrix genome_primer_matrix_12680.tsv \
        --lineage genome_lineage_12680.tsv \
        --out SupplementaryTables.xlsx \
        --sheet-name TableS1

    python3 merge_matrix_lineage.py \
        --matrix genome_primer_matrix_15658.tsv \
        --lineage genome_lineage_15658.tsv \
        --out SupplementaryTables.xlsx \
        --sheet-name TableS2
"""

import argparse
import os
import sys
import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def main():
    ap = argparse.ArgumentParser(
        description="Left-merge genome_primer_matrix.tsv with genome_lineage.tsv "
                    "on genome_id and export to Excel (Supplementary Table S1/S2 "
                    "column layout)."
    )
    ap.add_argument("--matrix", required=True,
                     help="Path to genome_primer_matrix.tsv (LEFT/primary table)")
    ap.add_argument("--lineage", required=True,
                     help="Path to genome_lineage.tsv (RIGHT table)")
    ap.add_argument("--out", default="merged_genome_taxonomy.xlsx",
                     help="Output .xlsx path")
    ap.add_argument("--sheet-name", default="genome_taxonomy_merged",
                     help="Sheet name, e.g. 'TableS1' or 'TableS2' to match the "
                          "manuscript's supplementary file")
    args = ap.parse_args()

    print(f"[merge] Reading {args.matrix} ...", file=sys.stderr)
    df_matrix = pd.read_csv(args.matrix, sep="\t", dtype=str)
    print(f"[merge]   {len(df_matrix)} rows, columns: {list(df_matrix.columns)}",
          file=sys.stderr)

    print(f"[merge] Reading {args.lineage} ...", file=sys.stderr)
    df_lineage = pd.read_csv(args.lineage, sep="\t", dtype=str)
    print(f"[merge]   {len(df_lineage)} rows, columns: {list(df_lineage.columns)}",
          file=sys.stderr)

    if "genome_id" not in df_matrix.columns:
        sys.exit(f"[merge] ERROR: 'genome_id' column not found in {args.matrix}")
    if "genome_id" not in df_lineage.columns:
        sys.exit(f"[merge] ERROR: 'genome_id' column not found in {args.lineage}")

    # Avoid duplicate genome_id rows in the right table causing a fan-out join
    dup_count = df_lineage["genome_id"].duplicated().sum()
    if dup_count > 0:
        print(f"[merge] WARNING: {dup_count} duplicated genome_id rows found in "
              f"{args.lineage}; keeping the first occurrence of each.",
              file=sys.stderr)
        df_lineage = df_lineage.drop_duplicates(subset="genome_id", keep="first")

    merged = df_matrix.merge(df_lineage, on="genome_id", how="left", validate="one_to_one")

    n_matched = merged["taxid"].notna().sum() if "taxid" in merged.columns else 0
    print(f"[merge] Merged table: {len(merged)} rows total, "
          f"{n_matched} rows matched a lineage record.", file=sys.stderr)

    # --- Write to Excel ---
    # If --out already exists, append/replace this one sheet in it instead
    # of overwriting the whole workbook, so TableS1 and TableS2 can end up
    # as two sheets in the same file across two separate runs.
    if os.path.exists(args.out):
        with pd.ExcelWriter(args.out, engine="openpyxl", mode="a",
                             if_sheet_exists="replace") as writer:
            merged.to_excel(writer, sheet_name=args.sheet_name, index=False)
    else:
        merged.to_excel(args.out, sheet_name=args.sheet_name, index=False)

    # Light formatting pass: bold header, frozen header row, reasonable
    # column widths.
    from openpyxl import load_workbook
    wb = load_workbook(args.out)
    ws = wb[args.sheet_name]

    header_font = Font(name="Arial", bold=True)
    body_font = Font(name="Arial")

    for col_idx, col_name in enumerate(merged.columns, start=1):
        col_letter = get_column_letter(col_idx)
        ws[f"{col_letter}1"].font = header_font
        ws[f"{col_letter}1"].alignment = Alignment(horizontal="center")

        max_len = max(
            [len(str(col_name))] +
            [len(str(v)) for v in merged[col_name].astype(str).values[:1000]]
        )
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font

    ws.freeze_panes = "A2"
    wb.save(args.out)

    print(f"[merge] Wrote {args.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
